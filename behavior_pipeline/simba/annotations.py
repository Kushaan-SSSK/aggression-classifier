# Kushaan Sharma
"""Turn the annotator's timestamp CSV into SimBA targets_inserted files.
Run as the `annotate` stage of the CLI; `--make-todo` writes the annotation sheet.
"""
from __future__ import annotations

from collections import defaultdict
from pathlib import Path

import pandas as pd

from ..config import LOG, PipelineConfig
from ..video.manifest import AnnotationRow, ScoredSpan, behavior_aliases, read_annotations, read_scored_spans, safe_video_name


def _video_fps(cfg: PipelineConfig) -> dict[str, float]:
    info = cfg.simba_project_folder / "logs" / "video_info.csv"
    if not info.exists():
        return {}
    df = pd.read_csv(info)
    return {str(r["Video"]): float(r["fps"]) for _, r in df.iterrows()}


def _project_classifiers(cfg: PipelineConfig) -> list[str]:
    ini_path = cfg.simba_config_path
    if not ini_path.exists():
        return []
    from .project import _read_ini

    ini = _read_ini(ini_path)
    sec = "SML settings"
    return [ini.get(sec, k) for k in ini[sec] if k.startswith("target_name_")] if ini.has_section(sec) else []


def mask_dir(cfg: PipelineConfig) -> Path:
    """Where the per-video label masks live (1 = label known for that frame and classifier)."""
    return cfg.simba_project_folder / "logs" / "label_masks"


def partial_classifiers(cfg: PipelineConfig) -> set[str]:
    """Classifiers whose labels are only known inside scored spans (the scored_span_groups)."""
    out: set[str] = set()
    for clfs in (cfg.get("annotations.scored_span_groups") or {}).values():
        out |= set(clfs)
    return out & set(cfg.classifiers)


def _frame_range(start_s: float, end_s: float, fps: float, n: int) -> tuple[int, int]:
    """Frames covered by [start_s, end_s); an event shorter than one frame still gets the frame it starts in."""
    a = max(int(round(start_s * fps)), 0)
    b = int(round(end_s * fps))
    if end_s > start_s and b <= a:
        a, b = int(start_s * fps), int(start_s * fps) + 1
    return a, min(b, n)


def build_targets(cfg: PipelineConfig, annotations_path: str | Path | None = None, features_dir: Path | None = None,
                  targets_dir: Path | None = None, scored_spans_path: str | Path | None = None) -> list[Path]:
    """Write one targets_inserted CSV per video (0/1 per classifier) plus a label mask, and return the target paths.

    The mask (logs/label_masks/<video>.csv) says which frames carry a known label per classifier: every
    frame for whole-video behaviours, only the scored spans plus labelled bouts for the partial ones.
    A bout ending well past the clip is flagged in logs/annotation_range_flags.csv and skipped.
    """
    annotations_path = Path(annotations_path) if annotations_path else cfg.path("paths.annotations")
    features_dir = features_dir or (cfg.simba_project_folder / "csv" / "features_extracted")
    targets_dir = targets_dir or (cfg.simba_project_folder / "csv" / "targets_inserted")
    targets_dir.mkdir(parents=True, exist_ok=True)
    masks_dir = mask_dir(cfg)
    masks_dir.mkdir(parents=True, exist_ok=True)
    rows = read_annotations(annotations_path, cfg)
    spans = read_scored_spans(cfg, scored_spans_path)
    classifiers = cfg.classifiers
    partial = partial_classifiers(cfg)
    # SimBA's trainer wants a column for every classifier registered in the project (e.g. Proximity)
    extra = [c for c in _project_classifiers(cfg) if c not in classifiers]
    none_token = str(cfg.get("annotations.none_token", "NONE")).upper()
    exclude_token = str(cfg.get("annotations.exclude_token", "EXCLUDE")).upper()
    exclude_clfs = [c for c in (cfg.get("annotations.exclude_classifiers") or ["Attack"]) if c in classifiers]
    fps_by_video = _video_fps(cfg)
    slack_s = float(cfg.get("annotations.end_slack_s", 30.0))   # a bout may run past the trimmed clip; a typo runs minutes past it

    by_video: dict[str, list[AnnotationRow]] = defaultdict(list)
    for r in rows:
        by_video[r.video_name].append(r)
    spans_by_video: dict[str, list[ScoredSpan]] = defaultdict(list)
    for s in spans:
        spans_by_video[s.video_name].append(s)

    feature_files = {p.stem: p for p in features_dir.glob("*.csv")}
    if not feature_files:
        raise FileNotFoundError(f"no feature files in {features_dir}; run the project stage first")

    videos = set(by_video)
    if cfg.get("annotations.include_unannotated_videos", False):
        videos |= set(feature_files)
    written: list[Path] = []
    summary = []
    flags = []
    for name in sorted(videos):
        if name not in feature_files:
            LOG.warning("%s: annotated but no features file; skipping", name)
            continue
        fps = fps_by_video.get(name, cfg.fps)
        df = pd.read_csv(feature_files[name], index_col=0)
        n = len(df)
        for clf in classifiers + extra:
            df[clf] = 0
        known = pd.DataFrame({clf: (0 if clf in partial else 1) for clf in classifiers}, index=range(n))
        known.index.name = "frame_idx"
        for s in spans_by_video.get(name, []):
            a, b = _frame_range(s.start_s, s.end_s, fps, n)
            for clf in s.classifiers:
                known.iloc[a:b, known.columns.get_loc(clf)] = 1
        counts = {clf: 0 for clf in classifiers}
        excluded = []
        for r in by_video.get(name, []):
            if r.behavior == none_token:
                continue
            if r.behavior == exclude_token:
                excluded.append(r)
                continue
            start, end = _frame_range(r.start_s, r.end_s, fps, n)
            if start >= n or r.end_s > n / fps + slack_s:
                LOG.warning("%s: %s %.1f-%.1f s lies outside the %.1f s clip; skipped, see logs/annotation_range_flags.csv",
                            name, r.behavior, r.start_s, r.end_s, n / fps)
                flags.append({"video": name, "behavior": r.behavior, "start_s": r.start_s, "end_s": r.end_s, "clip_s": round(n / fps, 1)})
                continue
            df.iloc[start:end, df.columns.get_loc(r.behavior)] = 1
            # a labelled bout is known even outside the scored spans (the annotator saw it)
            known.iloc[start:end, known.columns.get_loc(r.behavior)] = 1
            counts[r.behavior] += max(end - start, 0)
        # EXCLUDE: the annotator saw something there but it should neither train nor be scored; a labelled bout wins
        for r in excluded:
            a, b = _frame_range(r.start_s, r.end_s, fps, n)
            for clf in exclude_clfs:
                free = df.iloc[a:b][clf].to_numpy() == 0
                known.iloc[a:b, known.columns.get_loc(clf)] = known.iloc[a:b][clf].to_numpy() * ~free
        out = targets_dir / f"{name}.csv"
        df.to_csv(out)
        known.to_csv(masks_dir / f"{name}.csv")
        written.append(out)
        summary.append({"video": name, "frames": n, **{f"{c}_frames": counts[c] for c in classifiers},
                        **{f"{c}_known": int(known[c].sum()) for c in classifiers}})
        LOG.info("%s: targets written (%s)", name, ", ".join(f"{c}={counts[c]}/{int(known[c].sum())}" for c in classifiers))
    logs = cfg.simba_project_folder / "logs"
    if summary:
        pd.DataFrame(summary).to_csv(logs / "annotation_summary.csv", index=False)
    pd.DataFrame(flags, columns=["video", "behavior", "start_s", "end_s", "clip_s"]).to_csv(logs / "annotation_range_flags.csv", index=False)
    if flags:
        LOG.warning("%d annotation row(s) outside their clip, see logs/annotation_range_flags.csv", len(flags))
    return written


def read_mask(cfg: PipelineConfig, video_name: str, classifier: str, n: int | None = None) -> "pd.Series | None":
    """Known-label flags for one video and classifier, or None when no mask was written."""
    path = mask_dir(cfg) / f"{video_name}.csv"
    if not path.exists():
        return None
    m = pd.read_csv(path, index_col=0)
    if classifier not in m.columns:
        return None
    col = m[classifier].astype(int)
    return col.iloc[:n] if n is not None else col


def write_todo(cfg: PipelineConfig, out_path: str | Path | None = None) -> Path:
    """Write annotations_todo.csv with one TODO row per processed video."""
    import json

    out_path = Path(out_path) if out_path else cfg.path("paths.annotations").with_name("annotations_todo.csv")
    proc_dir = cfg.path("paths.processed_videos")
    rows = []
    for sidecar in sorted(proc_dir.glob("*.json")):
        meta = json.loads(sidecar.read_text(encoding="utf-8"))
        out = meta.get("output", {})
        fps = float(out.get("fps", cfg.fps))
        dur = float(out.get("duration_s", 0.0))
        rows.append({
            "video_name": meta.get("video_name", sidecar.stem), "behavior": "TODO", "start_s": "", "end_s": "", "annotator": "",
            "notes": (f"clip is 0-{dur:.0f} s; 0 = intruder entry (raw video {meta.get('intruder_entry_s', 0):.0f} s); "
                      f"seconds = frame counter / {fps:g}; behaviours: Attack, Social_investigation, or NONE if none occur"),
        })
    df = pd.DataFrame(rows, columns=["video_name", "behavior", "start_s", "end_s", "annotator", "notes"])
    out_path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(out_path, index=False)
    LOG.info("annotation todo sheet with %d videos -> %s", len(df), out_path)
    return out_path


def bouts_to_frames_table(rows: list[AnnotationRow], fps: float) -> pd.DataFrame:
    """Tabulate annotation rows as start/end frames."""
    recs = []
    for r in rows:
        if r.start_s is None:
            continue
        recs.append({"video_name": r.video_name, "behavior": r.behavior, "start_frame": int(round(r.start_s * fps)),
                     "end_frame": int(round(r.end_s * fps)), "duration_s": r.end_s - r.start_s})
    return pd.DataFrame(recs)


SKIP_SHEETS = ("explanation", "readme", "notes", "legend")


def _cell_seconds(value) -> float | None:
    """Seconds from a workbook cell: Excel parses '0:31' as 00:31:00, so a whole-minute time is read as m:ss."""
    import datetime as dt

    from ..video.manifest import parse_time_s

    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, dt.datetime):
        value = value.time()
    if isinstance(value, dt.time):
        if value.second == 0 and value.microsecond == 0:
            return value.hour * 60 + value.minute
        return value.hour * 3600 + value.minute * 60 + value.second + value.microsecond / 1e6
    if isinstance(value, (int, float)):
        return float(value)
    return parse_time_s(str(value))


def _split_remark(text: str, aliases: dict[str, str] | None = None) -> tuple[str, str]:
    """'Flank Investigation (low confidence)' -> ('Flank Investigation', 'low confidence'); a non-behaviour sentence -> ('', text).

    A label listed in the alias table is kept whole, however long or hedged it is.
    """
    import re

    text = text.strip()
    if aliases and text.lower() in aliases:
        return text, ""
    m = re.match(r"^([^(]+?)\s*\((.*?)\)?\s*$", text)
    beh, remark = (m.group(1).strip(), m.group(2).strip()) if m else (text, "")
    if aliases and beh.lower() in aliases:
        return beh, remark
    if len(beh.split()) > 4 or beh.upper().startswith(("NOT ", "THERE ", "STARTING", "SKIPPED", "CURRENTLY")):
        return "", text
    return beh, remark


def _merge_rows(new: pd.DataFrame, out_csv: Path) -> pd.DataFrame:
    """Rows of videos in `new` replace those in the existing CSV; other videos are kept."""
    if not out_csv.exists():
        return new
    import datetime as dt

    old = pd.read_csv(out_csv, dtype=str, keep_default_na=False)
    for col in new.columns:
        if col not in old.columns:
            old[col] = ""
    keys = {safe_video_name(v).lower() for v in new["video_name"].unique()}
    stale = old["video_name"].map(lambda v: safe_video_name(v).lower()).isin(keys)
    backup = out_csv.with_name(f"{out_csv.stem}.{dt.date.today().isoformat()}.bak.csv")
    if not backup.exists():
        old.to_csv(backup, index=False)
    replaced = sorted(old.loc[stale, "video_name"].unique())
    added = sorted(v for v in new["video_name"].unique() if safe_video_name(v).lower() not in
                   set(old["video_name"].map(lambda v: safe_video_name(v).lower())))
    LOG.info("merge into %s: replaced %d video(s) %s, added %d %s", out_csv.name, len(replaced), replaced, len(added), added)
    return pd.concat([old.loc[~stale, new.columns], new], ignore_index=True)


def _suggest_spans(sheet: str, segments: list[tuple[float, float]], exhaustive: bool, has_social: bool) -> None:
    if not segments or not has_social:
        return
    if not exhaustive:
        LOG.info("sheet %r: says its non-aggressive rows are not exhaustive; no scored span suggested", sheet)
        return
    txt = "; ".join(f"{a:.0f}-{b:.0f} s" for a, b in segments)
    LOG.info("sheet %r: scored stretches (raw seconds, for scored_spans.csv if the social rows are complete): %s", sheet, txt)


def source_fps_by_name(cfg: PipelineConfig) -> dict[str, float]:
    """Frame rate of each manifest video's raw recording, keyed by every spelling the annotator may use."""
    import json

    from ..video.manifest import read_video_manifest, video_name_lookup
    from ..video.preprocess import probe_video

    fps: dict[str, float] = {}
    if not cfg.path("paths.video_manifest").exists():
        return fps
    for e in read_video_manifest(cfg.path("paths.video_manifest"), cfg):
        side = cfg.path("paths.processed_videos") / f"{e.video_name}.json"
        rate = None
        if side.exists():
            rate = (json.loads(side.read_text(encoding="utf-8")).get("source") or {}).get("fps")
        if not rate and e.source_path.exists():
            rate = probe_video(e.source_path, cfg)["fps"]
        if rate:
            for key in (e.video_name, e.source_path.stem, e.source_path.name):
                fps[key.lower()] = float(rate)
    by_name = {name: fps[name.lower()] for name in {n for n in video_name_lookup(cfg).values()} if name.lower() in fps}
    return {spelling: by_name[name] for spelling, name in video_name_lookup(cfg).items() if name in by_name} | fps


def frame_timing_check(new: pd.DataFrame, old: pd.DataFrame, aliases: dict[str, str] | None = None, window_s: float = 5.0) -> dict:
    """Compare bouts converted from frame numbers with the same video's older second-based rows.

    For each new bout, the nearest older start of the same behaviour within window_s gives a time offset. The
    same is tried with the new times doubled and halved: frame numbers counted on a 30 fps copy of a 60 fps
    video come out at half the true time, so most bouts would line up only after doubling.
    """
    al = {k.lower(): v for k, v in (aliases or {}).items()}
    clf = lambda b: al.get(str(b).strip().lower(), str(b).strip())
    old = old[old["behavior"].astype(str).str.upper().ne("NONE")]
    old_by = {}
    for r in old.itertuples():
        try:
            old_by.setdefault((str(r.video_name).lower(), clf(r.behavior)), []).append(float(r.start_s))
        except (TypeError, ValueError):
            continue
    offsets, hits = [], {1.0: 0, 2.0: 0, 0.5: 0}
    for r in new.itertuples():
        starts = old_by.get((str(r.video_name).lower(), clf(r.behavior)))
        if not starts:
            continue
        for k in hits:
            d = min(abs(float(r.start_s) * k - s0) for s0 in starts)
            if d <= window_s:
                hits[k] += 1
                if k == 1.0:
                    offsets.append(min((float(r.start_s) - s0 for s0 in starts), key=abs))
    med = float(pd.Series(offsets).abs().median()) if offsets else None
    suspect = bool((hits[2.0] > 2 * max(hits[1.0], 1) or hits[0.5] > 2 * max(hits[1.0], 1)) or (med is not None and med > 1.5))
    return {"compared": len(new), "matched": hits[1.0], "matched_if_doubled": hits[2.0], "matched_if_halved": hits[0.5],
            "median_abs_offset_s": med, "suspect": suspect}


def _is_number(v) -> bool:
    if isinstance(v, bool) or v is None:
        return False
    if isinstance(v, (int, float)):
        return True
    try:
        float(str(v).strip())
        return True
    except ValueError:
        return False


def _header_cols(row) -> dict[str, int]:
    """Column positions of a workbook header row: start/end times, start/end frames, behaviour; {} if not a header."""
    cols: dict[str, int] = {}
    for i, c in enumerate(row):
        h = str(c).strip().lower() if c is not None else ""
        if h.startswith("start"):
            cols.setdefault("fstart" if "frame" in h else "start", i)
        elif h.startswith("end") or h.startswith("stop"):
            cols.setdefault("fend" if "frame" in h else "end", i)
        elif h.startswith("desc") or h.startswith("behav"):
            cols.setdefault("behavior", i)
    has_times = {"start", "end"} <= set(cols)
    has_frames = "fstart" in cols
    return cols if "behavior" in cols and (has_times or has_frames) else {}


def workbook_to_csv(xlsx: str | Path, out_csv: str | Path, time_reference: str = "raw", pad_end_s: float = 1.0,
                    annotator: str = "", merge: bool = True, aliases: dict[str, str] | None = None,
                    sheet_renames: dict[str, str] | None = None, span_classifiers: set[str] | None = None,
                    source_fps: dict[str, float] | None = None, frame_base: int = 0, frame_fps: float | None = None) -> pd.DataFrame:
    """One sheet per video, columns Start time / End time / Description, to the annotation CSV.

    Whole-second rows get pad_end_s added to the end; rows with fractions of a second are taken as written
    (a zero-length one lasts 0.5 s). Sheets whose start/end headers say "frame" hold frame numbers of the
    raw recording (first frame = frame_base, end frame inclusive), converted with source_fps[sheet video], or with
    frame_fps when the scorer's player counts at its own rate (the UT Box web player counts 30 fps on 60 fps videos).

    Times in the sheets are whole seconds, so a bout ending 'at 0:36' gets pad_end_s added and a
    single-second event becomes a one-second bout. With merge, videos already in the CSV are replaced
    and the others kept (a dated .bak copy is written first). Marker rows ("SKIPPED AHEAD",
    "Currently ended scoring") split each sheet into scored stretches, which are logged as
    suggestions for scored_spans.csv.
    """
    import re

    import openpyxl

    wb = openpyxl.load_workbook(str(xlsx), data_only=True)
    renames = {str(k).strip().lower(): str(v) for k, v in (sheet_renames or {}).items()}
    recs = []
    frame_videos: set[str] = set()
    for ws in wb.worksheets:
        if ws.title.strip().lower() in SKIP_SHEETS:
            continue
        rows = [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
        if not rows:
            continue
        video = renames.get(ws.title.strip().lower(), ws.title.strip())
        if video != ws.title.strip():
            LOG.info("sheet %r read as video %r", ws.title, video)
        # the header is the first row, or the second under a banner row ("Original | Original | Updated | Updated")
        hdr_row = next((k for k in (0, 1) if k < len(rows) and _header_cols(rows[k])), None)
        if hdr_row is None:
            LOG.warning("sheet %r: no Start/End/Description header, skipped", ws.title)
            continue
        header = [str(c).strip().lower() if c is not None else "" for c in rows[hdr_row]]
        cols = _header_cols(rows[hdr_row])
        frames_fps = None
        if "fstart" in cols:
            frames_fps = frame_fps or (source_fps or {}).get(video.lower()) or (source_fps or {}).get(ws.title.strip().lower())
            if not frames_fps:
                raise ValueError(f"sheet {ws.title!r} gives frame numbers but its recording's frame rate is unknown; "
                                 "add the video to the manifest first")
            LOG.info("sheet %r: frame numbers at %.3f fps (first frame %d)%s", ws.title, frames_fps, frame_base,
                     ", seconds where a row has no frames" if "start" in cols else "")
            frame_videos.add(video)
        extra = [i for i in range(len(header)) if i not in cols.values()]
        sheet_text = " ".join(str(c) for r in rows for c in r if c is not None).lower()
        exhaustive = "not exhaustive" not in sheet_text
        segments: list[tuple[float, float]] = []
        has_social = False
        seg_start: float | None = None
        seg_end: float | None = None
        for r in rows[hdr_row + 1:]:
            r = list(r) + [None] * (len(header) - len(r))
            first = str(r[cols.get("start", cols.get("fstart"))] or "").strip().lower()
            use_frames = "fstart" in cols and _is_number(r[cols["fstart"]])
            # "--" in the frame column: the scorer marked this event to be left out (excluded, not a negative)
            dashed = "fstart" in cols and str(r[cols["fstart"]] or "").strip().lower() in ("--", "-", "–", "—", "n/a", "na", "x", "exclude")
            # "merged above" in the frame column, or no frames and a note "grouping with attack below": the scorer folded
            # this row into a neighbouring bout, so its whole-second times must not come back as a bout of their own
            if "fstart" in cols and not use_frames and re.search(r"\bmerged\b|\bgroup(ed|ing)? with\b",
                                                                 " ".join(str(c) for c in r if c is not None), re.IGNORECASE):
                LOG.info("sheet %r: row folded into a neighbouring bout by the scorer, skipped: %s", ws.title,
                         " | ".join(str(c) for c in r if c is not None)[:80])
                continue
            if first.startswith(("no aggression", "none", "no attack")):
                # a fully scored video with nothing in it: every frame is a negative
                notes = "; ".join(str(c).strip() for c in r[1:] if c is not None and str(c).strip())
                recs.append({"video_name": video, "behavior": "NONE", "start_s": "", "end_s": "", "annotator": annotator,
                             "notes": notes, "time_reference": time_reference})
                continue
            try:
                if use_frames:
                    start = (float(r[cols["fstart"]]) - frame_base) / frames_fps
                elif "start" in cols:
                    start = _cell_seconds(r[cols["start"]])
                elif r[cols["fstart"]] is None:
                    start = None
                else:
                    raise ValueError(first)
            except Exception:
                # a remark in the time column ("SKIPPED AHEAD", "Currently ended scoring ...") marks a coverage gap
                LOG.info("sheet %r: remark row skipped: %s", ws.title, first[:60])
                if seg_start is not None:
                    segments.append((seg_start, seg_end))
                seg_start = seg_end = None
                continue
            if use_frames:
                fe = r[cols["fend"]] if "fend" in cols else None
                end = (float(fe) - frame_base + 1) / frames_fps if _is_number(fe) else None
            else:
                end = _cell_seconds(r[cols["end"]]) if "end" in cols else None
            beh = r[cols["behavior"]]
            if start is None or beh is None or not str(beh).strip():
                continue
            if use_frames:
                end = end if end is not None else start + 1 / frames_fps
            elif start % 1 or (end is not None and end % 1):
                end = max(end if end is not None else start, start + 0.5)
            else:
                end = (end if end is not None else start) + pad_end_s
            seg_start = start if seg_start is None else min(seg_start, start)
            seg_end = end if seg_end is None else max(seg_end, end)
            notes = "; ".join(str(r[i]).strip() for i in extra if r[i] is not None and str(r[i]).strip())
            # "Flank Investigation (low confidence)" -> behaviour + remark; a sentence that is not a behaviour is kept as a remark only
            beh, remark = _split_remark(str(beh), aliases)
            if remark:
                notes = "; ".join(x for x in (remark, notes) if x)
            if aliases and span_classifiers and aliases.get(beh.lower()) in span_classifiers:
                has_social = True
            if not beh:
                recs.append({"video_name": video, "behavior": "REMARK", "start_s": start, "end_s": end, "annotator": annotator,
                             "notes": notes, "time_reference": time_reference})
                continue
            label = str(beh).strip()
            if dashed and aliases and aliases.get(label.lower()) == "Attack":
                notes = "; ".join(x for x in (f"excluded by the scorer ('--' frames): {label}", notes) if x)
                label = "EXCLUDE"
            recs.append({"video_name": video, "behavior": label, "start_s": start, "end_s": end,
                         "annotator": annotator, "notes": notes, "time_reference": time_reference})
        if seg_start is not None:
            segments.append((seg_start, seg_end))
        _suggest_spans(ws.title, segments, exhaustive, has_social or span_classifiers is None)
    df = pd.DataFrame(recs, columns=["video_name", "behavior", "start_s", "end_s", "annotator", "notes", "time_reference"])
    out_csv = Path(out_csv)
    out_csv.parent.mkdir(parents=True, exist_ok=True)
    LOG.info("%d bouts from %d sheet(s) in %s", len(df), df["video_name"].nunique(), Path(xlsx).name)
    if frame_videos and out_csv.exists():
        chk = frame_timing_check(df[df["video_name"].isin(frame_videos)], pd.read_csv(out_csv, dtype={"video_name": str}), aliases)
        (LOG.warning if chk["suspect"] else LOG.info)(
            "frame-number timing vs the older second-based rows: %s%s", chk,
            " -- CHECK the frame numbering (a 30 fps count of a 60 fps video halves every time)" if chk["suspect"] else "")
    if merge:
        df = _merge_rows(df, out_csv)
    df.to_csv(out_csv, index=False)
    LOG.info("%d rows -> %s", len(df), out_csv)
    return df


def _yes(value) -> bool | None:
    """Review-sheet answer: True / False, or None when blank or not understood."""
    t = " ".join(str(value or "").strip().lower().split())
    if t in ("y", "yes", "x", "true", "1", "attack", "most likely yes", "likely yes"):
        return True
    if t in ("n", "no", "false", "0", "not attack", "no attack", "most likely not", "likely not", "most likely not enough"):
        return False
    return None


def note_events(text: str) -> tuple[list[tuple[float, float]], list[float]]:
    """(start, end) pairs written in a note ("second bite starts at 0:17.28 and ends at 0:17.32") and unpaired times."""
    import re

    from ..video.manifest import parse_time_s

    times = [parse_time_s(t) for t in re.findall(r"(\d+:\d{2}(?:\.\d+)?)", str(text or ""))]
    pairs, lone, k = [], [], 0
    while k < len(times):
        if k + 1 < len(times) and 0 < times[k + 1] - times[k] <= 5:
            pairs.append((times[k], times[k + 1]))
            k += 2
        else:
            lone.append(times[k])
            k += 1
    return pairs, lone


def import_review_sheet(cfg: PipelineConfig, xlsx: str | Path, bouts_csv: str | Path | None = None,
                        annotations_path: str | Path | None = None, annotator: str = "NR",
                        attack_rows: set[int] | None = None, not_attack_rows: set[int] | None = None) -> pd.DataFrame:
    """Apply an annotator's answers on a review sheet (missed bouts or false-alarm candidates) to the annotation CSV.

    The sheet's first column is the bout number in bouts.csv (written next to it when the review package was made:
    video, raw_start_s, raw_end_s). New times come as raw-video frame numbers ("Start frame", first frame =
    annotations.frame_base) or as times ("New start (m:ss.s)"); further events written in the notes as
    "starts at 0:17.28 and ends at 0:17.32" are added too. attack_rows / not_attack_rows override hedged answers.
    missed-bout sheet  - the Attack rows overlapping the reviewed bout (or any new event) are replaced by the new
                         events (attack) or by a not-attack row, keeping every other row of the video;
    false-alarm sheet  - attack adds the missed events; not attack adds the named behaviour when it maps to a classifier.
    """
    import datetime as dt
    import shutil

    import openpyxl

    from ..video.manifest import behavior_aliases, entry_times, parse_time_s, resolve_video_name, video_name_lookup

    xlsx = Path(xlsx)
    bouts = pd.read_csv(Path(bouts_csv) if bouts_csv else xlsx.parent / "bouts.csv").set_index("n")
    ann_path = Path(annotations_path) if annotations_path else cfg.path("paths.annotations")
    ann = pd.read_csv(ann_path, dtype=str, keep_default_na=False)
    aliases = {k.lower(): v for k, v in behavior_aliases(cfg).items()}
    fps_by = source_fps_by_name(cfg)
    base = int(cfg.get("annotations.frame_base", 0))
    lookup, entries = video_name_lookup(cfg), entry_times(cfg)
    attack_rows, not_attack_rows = set(attack_rows or ()), set(not_attack_rows or ())

    ws = openpyxl.load_workbook(str(xlsx), data_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    head = [str(c or "").strip().lower() for c in rows[0]]

    def col(*keys):
        return next((i for i, h in enumerate(head) if any(h.startswith(k) for k in keys)), None)

    def cell(r, c):
        return r[c] if c is not None and c < len(r) else None

    c_ans = col("is it an attack", "missed attack")
    c_fs, c_fe = col("start frame"), col("end frame")
    c_ts, c_te = col("new start"), col("new end")
    c_beh, c_note = col("behaviour", "behavior", "what is it"), col("notes")
    missed_sheet = head[c_ans].startswith("is it") if c_ans is not None else True

    def seconds(v) -> float | None:
        if v is None or str(v).strip() == "":
            return None
        try:
            return parse_time_s(v.strip()) if isinstance(v, str) else _cell_seconds(v)
        except Exception:
            return None

    def raw_window(r) -> tuple[float, float]:
        s, e = float(r["start_s"]), float(r["end_s"])
        if str(r.get("time_reference", "raw")) == "processed":
            shift = entries.get(resolve_video_name(str(r["video_name"]), lookup), 0.0)
            s, e = s + shift, e + shift
        return s, e

    stamp = dt.date.today().isoformat()
    added, removed, log = [], [], []
    for r in rows[1:]:
        if not r or r[0] is None or int(r[0]) not in bouts.index:
            continue
        n = int(r[0])
        b = bouts.loc[n]
        video = str(b["video"])
        fps = fps_by.get(video.lower()) or float(b.get("source_fps", 60.0))
        orig = (float(b["raw_start_s"]), float(b["raw_end_s"]))
        events: list[tuple[float, float]] = []
        if str(cell(r, c_fs) if cell(r, c_fs) is not None else "").strip().isdigit():
            f0 = float(cell(r, c_fs))
            f1 = float(cell(r, c_fe)) if str(cell(r, c_fe) or "").strip().isdigit() else f0
            events.append(((f0 - base) / fps, (f1 - base + 1) / fps))
        else:
            t0, t1 = seconds(cell(r, c_ts)), seconds(cell(r, c_te))
            if t0 is not None:
                events.append((t0, t1 if t1 is not None and t1 > t0 else t0 + 1 / fps))
        note = str(cell(r, c_note) or "").strip()
        extra, lone = note_events(note)
        events += extra
        ans = True if n in attack_rows else False if n in not_attack_rows else _yes(cell(r, c_ans))
        if ans is None:
            if not events:
                continue                               # not answered
            ans = True
        beh = str(cell(r, c_beh) or "").strip()
        tag = f"review {xlsx.stem} #{n} ({stamp})" + (f"; {note}" if note else "")
        if lone:
            LOG.warning("review #%d %s: time(s) %s in the notes have no start/end partner; not imported", n, video, lone)
        if missed_sheet:
            # the reviewed bout's labels are replaced; a new event elsewhere only replaces short coarse labels it refines
            # (a whole-second "Bite 0:17-0:18"), never a separately labelled longer bout it happens to touch
            for i, a in ann[ann["video_name"].str.lower() == video.lower()].iterrows():
                if aliases.get(a["behavior"].strip().lower()) != "Attack":
                    continue
                s0, e0 = raw_window(a)
                near = lambda w: s0 < w[1] + 0.5 and e0 > w[0] - 0.5
                if near(orig) or (e0 - s0 <= 2.0 and any(near(w) for w in events)):
                    removed.append(i)
        if ans:
            label = beh if aliases.get(beh.lower()) == "Attack" else "Attack"
            events = events or [orig]
        else:
            label = beh if beh and aliases.get(beh.lower()) not in (None, "Attack") else ("not attack" if missed_sheet else "")
            if not label:
                log.append({"n": n, "video": video, "answer": cell(r, c_ans), "action": "no change (not attack, behaviour not a classifier)"})
                continue
            events = events[:1] or [orig]
        for start, end in events:
            if abs(start - orig[0]) > 4:
                LOG.warning("review #%d %s: event at %.2f s is %.1f s from the reviewed bout (%.1f s)", n, video, start, start - orig[0], orig[0])
            added.append({"video_name": video, "behavior": label, "start_s": round(start, 3), "end_s": round(end, 3),
                          "annotator": annotator, "notes": tag, "time_reference": "raw"})
        log.append({"n": n, "video": video, "answer": cell(r, c_ans),
                    "action": f"{label}: " + ", ".join(f"{a:.2f}-{e:.2f}" for a, e in events)})
    if not added and not removed:
        LOG.info("%s: no answered rows yet", xlsx.name)
        return pd.DataFrame(log)
    bak = ann_path.with_name(f"{ann_path.stem}.pre-review-{stamp}.bak.csv")
    shutil.copy(ann_path, bak)
    out = pd.concat([ann.drop(index=sorted(set(removed))), pd.DataFrame(added).astype(str)], ignore_index=True)
    out.to_csv(ann_path, index=False)
    LOG.info("review %s: %d row(s) removed, %d added (backup %s)", xlsx.name, len(set(removed)), len(added), bak.name)
    return pd.DataFrame(log)
